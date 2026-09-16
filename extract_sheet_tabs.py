"""
Split the team Google Sheet (downloaded as .xlsx) into the two per-tab CSVs the
Google Ads + Demand builders expect, with the EXACT date formats each builder
parses. Getting the date format wrong silently drops every row, so this is
encoded once here instead of by hand.

WHY THIS EXISTS
---------------
google_build.py / demand_build.py each read ONE sheet tab exported to CSV:
  * gg_raw_creative -> google_build.py   (cost / impr / clicks for Google Ads)
  * raw_retention   -> demand_build.py   (lead / DAU / cost for the Demand phase)

The Drive API can only export a Google Sheet's FIRST tab as CSV, so to get a
specific tab we download the whole workbook as .xlsx and pull the tab out here.

DATE FORMATS (do not change without checking the builder):
  * gg_raw_creative "Day"  -> "%Y-%m-%d"   (google_build uses the string as-is,
                                            and it must equal Airbridge's
                                            CAST(Event_Date AS STRING) day key)
  * raw_retention  "date"  -> "%-m/%-d/%Y" (demand_build._iso() splits on "/" to
                                            turn M/D/YYYY into YYYY-MM-DD)

The gg_raw_creative tab has a junk banner row above the real header, so we locate
the header row by its first cell == "Day". raw_retention's header is row 1 and
carries duplicate account_name / Channel columns; we preserve column order so
demand_build's first-occurrence-wins column lookup behaves identically.

USAGE
-----
  python extract_sheet_tabs.py <workbook.xlsx> <out_dir>
    -> writes <out_dir>/gg_raw_creative.csv and <out_dir>/raw_retention.csv

Requires openpyxl (not in requirements.txt because CI never runs this; the
manual/scheduled refresh installs it into a throwaway venv).
"""

import csv
import datetime
import pathlib
import sys

# tab name -> (output csv name, date format for date-typed cells, header anchor)
# header anchor = the value of the header row's first cell (None => header is row 0).
TABS = {
    "gg_raw_creative": ("gg_raw_creative.csv", "%Y-%m-%d", "Day"),
    "raw_retention":   ("raw_retention.csv",   "%-m/%-d/%Y", None),
}


def fmt(value, date_fmt):
    """Render an openpyxl cell value as a clean CSV string.

    Dates use the per-tab format; integer-valued floats lose the ".0" (so IDs
    like 190329934796.0 don't corrupt the campaign/ad-group join key)."""
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime(date_fmt)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, int):
        return str(value)
    return str(value)


def export(ws, out_path, date_fmt, header_anchor):
    rows = list(ws.iter_rows(values_only=True))
    if header_anchor is not None:
        hi = next(i for i, r in enumerate(rows)
                  if r and str(r[0]).strip() == header_anchor)
    else:
        hi = 0
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([fmt(c, date_fmt) for c in rows[hi]])
        n = 0
        for r in rows[hi + 1:]:
            if not r or all(c in (None, "") for c in r):
                continue
            w.writerow([fmt(c, date_fmt) for c in r])
            n += 1
    return n


def main():
    if len(sys.argv) != 3:
        sys.exit("usage: python extract_sheet_tabs.py <workbook.xlsx> <out_dir>")
    import openpyxl  # imported lazily so `-h` works without the dep installed
    xlsx, out_dir = sys.argv[1], pathlib.Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    for tab, (out_name, date_fmt, anchor) in TABS.items():
        if tab not in wb.sheetnames:
            sys.exit(f"tab '{tab}' not found in workbook; tabs = {wb.sheetnames}")
        n = export(wb[tab], out_dir / out_name, date_fmt, anchor)
        print(f"{tab}: {n} rows -> {out_dir / out_name}")


if __name__ == "__main__":
    main()
